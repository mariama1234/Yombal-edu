"""
Exports Excel et PDF (« Exports Excel & PDF » du Module 7).

Le `.xlsx` passe par `openpyxl`, déjà installé pour l'import des listes
d'élèves. Le PDF est écrit ici, à la main.

Pourquoi écrire le PDF plutôt qu'installer ReportLab ou WeasyPrint : les
documents produits sont des tableaux et des reçus — du texte cadré, pas de la
mise en page riche. Une bibliothèque de rendu, c'est une dépendance lourde
(souvent des binaires système pour WeasyPrint) à installer sur chaque poste
d'école pour imprimer un reçu de 300 F. Le module ne connaît que les polices
standard PDF, présentes dans tous les lecteurs : rien à embarquer.

Le générateur couvre : titres, champs libellé/valeur, tableaux paginés avec
en-têtes répétés, totaux et pied de page numéroté.
"""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------

_ENTETE_FOND = PatternFill("solid", start_color="1F3A5F")
_ENTETE_POLICE = Font(bold=True, color="FFFFFF")


def classeur_xlsx(feuilles):
    """
    `feuilles` : liste de (nom, entetes, lignes).

    Les colonnes sont dimensionnées d'après le contenu — un export qu'il faut
    élargir colonne par colonne avant de le lire n'est pas un export.
    """
    classeur = Workbook()
    classeur.remove(classeur.active)

    for nom, entetes, lignes in feuilles:
        feuille = classeur.create_sheet(title=str(nom)[:31])
        feuille.append(list(entetes))
        for cellule in feuille[1]:
            cellule.fill = _ENTETE_FOND
            cellule.font = _ENTETE_POLICE
            cellule.alignment = Alignment(vertical="center")
        for ligne in lignes:
            feuille.append(list(ligne))

        feuille.freeze_panes = "A2"
        for index, entete in enumerate(entetes, start=1):
            largeur = max(
                [len(str(entete))] + [len(str(l[index - 1])) for l in lignes if len(l) >= index]
                or [len(str(entete))]
            )
            feuille.column_dimensions[get_column_letter(index)].width = min(max(largeur + 3, 11), 45)

    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

# Largeurs Helvetica (millièmes de cadratin), pour les caractères ASCII.
# Elles servent à centrer, aligner à droite et tronquer proprement : sans
# elles, un montant long déborderait sur la colonne voisine.
_LARGEURS = {
    " ": 278, "!": 278, '"': 355, "#": 556, "$": 556, "%": 889, "&": 667, "'": 191,
    "(": 333, ")": 333, "*": 389, "+": 584, ",": 278, "-": 333, ".": 278, "/": 278,
    ":": 278, ";": 278, "<": 584, "=": 584, ">": 584, "?": 556, "@": 1015,
    "[": 278, "\\": 278, "]": 278, "^": 469, "_": 556, "`": 333,
    "{": 334, "|": 260, "}": 334, "~": 584,
    "A": 667, "B": 667, "C": 722, "D": 722, "E": 667, "F": 611, "G": 778, "H": 722,
    "I": 278, "J": 500, "K": 667, "L": 556, "M": 833, "N": 722, "O": 778, "P": 667,
    "Q": 778, "R": 722, "S": 667, "T": 611, "U": 722, "V": 667, "W": 944, "X": 667,
    "Y": 667, "Z": 611,
    "a": 556, "b": 556, "c": 500, "d": 556, "e": 556, "f": 278, "g": 556, "h": 556,
    "i": 222, "j": 222, "k": 500, "l": 222, "m": 833, "n": 556, "o": 556, "p": 556,
    "q": 556, "r": 333, "s": 500, "t": 278, "u": 556, "v": 500, "w": 722, "x": 500,
    "y": 500, "z": 500,
}
_LARGEURS.update({chiffre: 556 for chiffre in "0123456789"})

# En Helvetica, un caractère accentué a la largeur de sa lettre de base.
_BASE = {
    "à": "a", "â": "a", "ä": "a", "á": "a", "ã": "a", "å": "a", "ç": "c",
    "è": "e", "é": "e", "ê": "e", "ë": "e", "î": "i", "ï": "i", "í": "i", "ì": "i",
    "ô": "o", "ö": "o", "ó": "o", "ò": "o", "õ": "o", "ù": "u", "û": "u", "ü": "u",
    "ú": "u", "ñ": "n", "ý": "y", "ÿ": "y",
    "À": "A", "Â": "A", "Ä": "A", "Á": "A", "Ç": "C", "È": "E", "É": "E", "Ê": "E",
    "Ë": "E", "Î": "I", "Ï": "I", "Ô": "O", "Ö": "O", "Ù": "U", "Û": "U", "Ü": "U",
    "Ñ": "N",
    "€": "E", "°": "o", "’": "'", "«": '"', "»": '"', "–": "-", "—": "-", "…": ".",
}

# Le gras d'Helvetica est un peu plus large que le romain ; le rapport suffit
# pour aligner et tronquer, on ne cherche pas la justification typographique.
_FACTEUR_GRAS = 1.05

A4 = (595.28, 841.89)
MARGE = 42.0
_INTERLIGNE = 13.0


def largeur_texte(texte, taille, gras=False):
    total = 0
    for caractere in str(texte):
        base = _BASE.get(caractere, caractere)
        total += _LARGEURS.get(base, 500)
    largeur = total / 1000.0 * taille
    return largeur * _FACTEUR_GRAS if gras else largeur


def tronquer(texte, largeur_max, taille, gras=False):
    texte = str(texte)
    if largeur_texte(texte, taille, gras) <= largeur_max:
        return texte
    coupe = texte
    while coupe and largeur_texte(coupe + "…", taille, gras) > largeur_max:
        coupe = coupe[:-1]
    return (coupe + "…") if coupe else ""


def _echapper(texte):
    """Passe le texte en WinAnsi et neutralise les caractères de syntaxe PDF."""
    brut = str(texte).encode("cp1252", "replace")
    return brut.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


class _Page:
    """Un flux de contenu en construction."""

    def __init__(self):
        self.operations = []

    def texte(self, x, y, contenu, taille=10, gras=False, gris=0.0):
        police = b"/F2" if gras else b"/F1"
        self.operations.append(
            b"BT %.3f g %s %.2f Tf %.2f %.2f Td (%s) Tj ET"
            % (gris, police, taille, x, y, _echapper(contenu))
        )

    def rectangle(self, x, y, largeur, hauteur, gris=0.85):
        self.operations.append(
            b"%.3f g %.2f %.2f %.2f %.2f re f" % (gris, x, y, largeur, hauteur))

    def flux(self):
        return b"\n".join(self.operations)


class DocumentPDF:
    """
    Rédaction séquentielle d'un document : on descend, le saut de page est
    automatique et les en-têtes de tableau sont réimprimés en haut de page.
    """

    def __init__(self, titre, sous_titre=None, pied=None, paysage=False):
        self.titre = titre
        self.sous_titre = sous_titre
        self.pied = pied or "YombalEdu"
        self.largeur, self.hauteur = (A4[1], A4[0]) if paysage else A4
        self.pages = []
        self._colonnes = None
        self._nouvelle_page()

    # -- mise en page ------------------------------------------------------
    @property
    def page(self):
        return self.pages[-1]

    def _nouvelle_page(self):
        self.pages.append(_Page())
        self.y = self.hauteur - MARGE
        self._entete()

    def _entete(self):
        premiere = len(self.pages) == 1
        self.page.texte(MARGE, self.y - 12, self.titre, taille=16 if premiere else 11,
                        gras=True, gris=0.11)
        self.y -= 20 if premiere else 16
        if self.sous_titre and premiere:
            self.page.texte(MARGE, self.y - 10, self.sous_titre, taille=9.5, gris=0.42)
            self.y -= 16
        self.page.rectangle(MARGE, self.y - 4, self.largeur - 2 * MARGE, 0.8, gris=0.75)
        self.y -= 18

    def _place(self, hauteur_necessaire):
        """Ouvre une page si le bloc ne tient plus, et réimprime les en-têtes."""
        if self.y - hauteur_necessaire >= MARGE + 24:
            return
        self._nouvelle_page()
        if self._colonnes:
            self._entetes_tableau(self._colonnes)

    def espace(self, hauteur=10):
        self.y -= hauteur

    # -- blocs -------------------------------------------------------------
    def section(self, texte):
        self._place(26)
        self.espace(4)
        self.page.texte(MARGE, self.y - 10, texte, taille=11, gras=True, gris=0.15)
        self.y -= 22

    def paragraphe(self, texte, taille=9.5, gris=0.3):
        for ligne in self._envelopper(texte, self.largeur - 2 * MARGE, taille):
            self._place(_INTERLIGNE)
            self.page.texte(MARGE, self.y - taille, ligne, taille=taille, gris=gris)
            self.y -= _INTERLIGNE

    def _envelopper(self, texte, largeur_max, taille):
        lignes, courante = [], ""
        for mot in str(texte).split():
            essai = f"{courante} {mot}".strip()
            if largeur_texte(essai, taille) <= largeur_max or not courante:
                courante = essai
            else:
                lignes.append(courante)
                courante = mot
        return lignes + [courante] if courante else lignes

    def champs(self, valeurs, largeur_libelle=140):
        """Bloc « libellé : valeur », utilisé par le reçu."""
        for libelle, valeur in valeurs:
            self._place(16)
            self.page.texte(MARGE, self.y - 10, libelle, taille=9.5, gris=0.45)
            self.page.texte(MARGE + largeur_libelle, self.y - 10, valeur, taille=10,
                            gras=True, gris=0.11)
            self.y -= 16

    def encadre(self, libelle, valeur, gris_fond=0.94):
        self._place(34)
        largeur = self.largeur - 2 * MARGE
        self.page.rectangle(MARGE, self.y - 30, largeur, 30, gris=gris_fond)
        self.page.texte(MARGE + 10, self.y - 12, libelle, taille=9, gris=0.4)
        largeur_valeur = largeur_texte(valeur, 13, gras=True)
        self.page.texte(MARGE + largeur - 10 - largeur_valeur, self.y - 21, valeur,
                        taille=13, gras=True, gris=0.11)
        self.y -= 38

    # -- tableaux ----------------------------------------------------------
    def tableau(self, colonnes, lignes, total=None):
        """
        `colonnes` : liste de (titre, largeur, alignement) — alignement 'g'
        (gauche) ou 'd' (droite, pour les montants).
        """
        colonnes = self._ajuster(colonnes)
        self._colonnes = colonnes
        self._place(40)
        self._entetes_tableau(colonnes)

        for index, ligne in enumerate(lignes):
            self._place(_INTERLIGNE + 4)
            if index % 2 == 1:
                self.page.rectangle(MARGE, self.y - 4, self.largeur - 2 * MARGE,
                                    _INTERLIGNE + 3, gris=0.965)
            self._ligne_tableau(colonnes, ligne, taille=9)
            self.y -= _INTERLIGNE + 3

        if total is not None:
            self._place(_INTERLIGNE + 8)
            self.page.rectangle(MARGE, self.y + _INTERLIGNE - 2,
                                self.largeur - 2 * MARGE, 0.8, gris=0.75)
            self._ligne_tableau(colonnes, total, taille=9.5, gras=True)
            self.y -= _INTERLIGNE + 6

        self._colonnes = None
        self.espace(6)

    def _ajuster(self, colonnes):
        """Répartit la largeur disponible au prorata des largeurs demandées."""
        disponible = self.largeur - 2 * MARGE
        demande = sum(c[1] for c in colonnes) or 1
        facteur = disponible / demande
        return [(titre, largeur * facteur, align) for titre, largeur, align in colonnes]

    def _entetes_tableau(self, colonnes):
        self.page.rectangle(MARGE, self.y - 5, self.largeur - 2 * MARGE,
                            _INTERLIGNE + 4, gris=0.90)
        self._ligne_tableau(colonnes, [c[0] for c in colonnes], taille=8.5, gras=True,
                            gris=0.2)
        self.y -= _INTERLIGNE + 5

    def _ligne_tableau(self, colonnes, cellules, taille, gras=False, gris=0.15):
        x = MARGE
        for (_, largeur, align), valeur in zip(colonnes, list(cellules) + [""] * len(colonnes)):
            texte = tronquer(valeur, largeur - 10, taille, gras)
            if align == "d":
                depart = x + largeur - 6 - largeur_texte(texte, taille, gras)
            else:
                depart = x + 4
            self.page.texte(depart, self.y, texte, taille=taille, gras=gras, gris=gris)
            x += largeur

    # -- sortie ------------------------------------------------------------
    def _pieds_de_page(self):
        horodatage = datetime.now().strftime("%d/%m/%Y à %H:%M")
        total = len(self.pages)
        for numero, page in enumerate(self.pages, start=1):
            page.rectangle(MARGE, MARGE + 14, self.largeur - 2 * MARGE, 0.6, gris=0.85)
            page.texte(MARGE, MARGE, f"{self.pied} — édité le {horodatage}",
                       taille=8, gris=0.5)
            mention = f"Page {numero} / {total}"
            page.texte(self.largeur - MARGE - largeur_texte(mention, 8), MARGE,
                       mention, taille=8, gris=0.5)

    def rendu(self):
        """Assemble le fichier PDF (table xref calculée sur les offsets réels)."""
        self._pieds_de_page()

        objets = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            None,  # 2 : arbre des pages, rempli une fois les numéros connus
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>",
        ]

        numeros_pages = []
        for page in self.pages:
            flux = page.flux()
            numero_contenu = len(objets) + 2
            numero_page = len(objets) + 1
            numeros_pages.append(numero_page)
            objets.append(
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 %.2f %.2f] "
                b"/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> /Contents %d 0 R >>"
                % (self.largeur, self.hauteur, numero_contenu)
            )
            objets.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(flux), flux))

        kids = b" ".join(b"%d 0 R" % n for n in numeros_pages)
        objets[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(numeros_pages))

        sortie = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
        offsets = []
        for index, corps in enumerate(objets, start=1):
            offsets.append(len(sortie))
            sortie += b"%d 0 obj\n" % index + corps + b"\nendobj\n"

        depart_xref = len(sortie)
        sortie += b"xref\n0 %d\n" % (len(objets) + 1)
        sortie += b"0000000000 65535 f \n"
        for offset in offsets:
            sortie += b"%010d 00000 n \n" % offset
        sortie += (b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
                   % (len(objets) + 1, depart_xref))
        return bytes(sortie)
